"""Professional Excel spreadsheet generator using openpyxl."""

import io
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from assistant_core.artifacts.schemas import WorkbookSpec

# Design tokens
HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color="0F172A")

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)

TOTAL_BORDER = Border(
    top=Side(style="thin", color="0F172A"),
    bottom=Side(style="double", color="0F172A"),
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
)

NUMBER_FORMATS = {
    "currency": "$#,##0.00",
    "percent": "0.0%",
    "integer": "#,##0",
    "float": "#,##0.00",
    "date": "yyyy-mm-dd",
}


class XlsxGenerator:
    """Renders structured WorkbookSpec into polished .xlsx bytes."""

    @staticmethod
    def generate(spec: WorkbookSpec) -> bytes:
        wb = openpyxl.Workbook()
        # Remove default sheet
        if wb.active:
            wb.remove(wb.active)

        for sheet_spec in spec.sheets:
            ws = wb.create_sheet(title=sheet_spec.name[:31])
            ws.views.sheetView[0].showGridLines = True

            # 1. Write headers
            if sheet_spec.headers:
                ws.append(sheet_spec.headers)
                for col_idx in range(1, len(sheet_spec.headers) + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.fill = HEADER_FILL
                    cell.font = HEADER_FONT
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                    cell.border = THIN_BORDER
                ws.row_dimensions[1].height = 28

            # 2. Write data rows
            current_row = 2 if sheet_spec.headers else 1
            start_data_row = current_row

            for r_idx, row_data in enumerate(sheet_spec.rows):
                ws.append(row_data)
                fill = ZEBRA_FILL if r_idx % 2 == 1 else WHITE_FILL

                for c_idx, val in enumerate(row_data):
                    col_num = c_idx + 1
                    cell = ws.cell(row=current_row, column=col_num)
                    cell.fill = fill
                    cell.border = THIN_BORDER
                    cell.font = Font(name="Calibri", size=10.5)

                    # Apply column type formatting
                    col_type = None
                    if sheet_spec.column_types and c_idx < len(sheet_spec.column_types):
                        col_type = sheet_spec.column_types[c_idx]

                    if col_type in NUMBER_FORMATS:
                        cell.number_format = NUMBER_FORMATS[col_type]
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif isinstance(val, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

                ws.row_dimensions[current_row].height = 20
                current_row += 1

            end_data_row = current_row - 1

            # 3. Add totals row if requested
            if sheet_spec.totals_row and sheet_spec.rows and sheet_spec.headers:
                totals: list[Any] = []
                for c_idx in range(len(sheet_spec.headers)):
                    col_type = None
                    if sheet_spec.column_types and c_idx < len(sheet_spec.column_types):
                        col_type = sheet_spec.column_types[c_idx]

                    if c_idx == 0:
                        totals.append("Total")
                    elif col_type in ("currency", "integer", "float") or any(
                        isinstance(r[c_idx], (int, float))
                        for r in sheet_spec.rows
                        if len(r) > c_idx
                    ):
                        col_letter = get_column_letter(c_idx + 1)
                        totals.append(
                            f"=SUM({col_letter}{start_data_row}:{col_letter}{end_data_row})"
                        )
                    else:
                        totals.append("")

                ws.append(totals)
                for c_idx in range(len(totals)):
                    col_num = c_idx + 1
                    cell = ws.cell(row=current_row, column=col_num)
                    cell.font = TOTAL_FONT
                    cell.border = TOTAL_BORDER
                    cell.fill = WHITE_FILL

                    col_type = None
                    if sheet_spec.column_types and c_idx < len(sheet_spec.column_types):
                        col_type = sheet_spec.column_types[c_idx]
                    if col_type in NUMBER_FORMATS:
                        cell.number_format = NUMBER_FORMATS[col_type]

                    if c_idx == 0:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="right", vertical="center")

                ws.row_dimensions[current_row].height = 24

            # 4. Auto-fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(int(col[0].column))
                for cell in col:
                    val_str = str(cell.value or "")
                    if (
                        cell.number_format
                        and "0.00" in cell.number_format
                        and isinstance(cell.value, (int, float))
                    ):
                        val_str = f"${cell.value:,.2f}"
                    max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
