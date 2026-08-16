"""Professional PowerPoint presentation (.pptx) generator using python-pptx."""

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from assistant_core.artifacts.schemas import PresentationSpec, SlideSpec

# Brand Palette
DARK_BG = RGBColor(15, 23, 42)  # Slate 900
LIGHT_BG = RGBColor(248, 250, 252)  # Slate 50
CARD_BG = RGBColor(241, 245, 249)  # Slate 100
CARD_BORDER = RGBColor(203, 213, 225)  # Slate 300
PRIMARY_TEXT = RGBColor(15, 23, 42)  # Slate 900
SECONDARY_TEXT = RGBColor(71, 85, 105)  # Slate 600
MUTED_TEXT = RGBColor(148, 163, 184)  # Slate 400
ACCENT_BLUE = RGBColor(59, 130, 246)  # Blue 500
WHITE = RGBColor(255, 255, 255)


class PptxGenerator:
    """Renders structured PresentationSpec into polished 16:9 .pptx bytes."""

    @staticmethod
    def generate(spec: PresentationSpec) -> bytes:
        prs = Presentation()
        # Set 16:9 Widescreen dimensions
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_slide_layout = prs.slide_layouts[6]  # Blank slide

        # 1. Title Slide
        title_slide = prs.slides.add_slide(blank_slide_layout)
        PptxGenerator._render_title_slide(title_slide, spec.title, spec.subtitle)

        # 2. Content Slides
        for s_idx, slide_spec in enumerate(spec.slides):
            slide = prs.slides.add_slide(blank_slide_layout)
            PptxGenerator._render_header(slide, slide_spec.title, slide_spec.subtitle, slide_num=s_idx + 2)

            if slide_spec.layout == "cards" and slide_spec.cards:
                PptxGenerator._render_cards_layout(slide, slide_spec)
            elif slide_spec.layout == "comparison" and (slide_spec.left_column or slide_spec.right_column):
                PptxGenerator._render_comparison_layout(slide, slide_spec)
            elif slide_spec.layout == "quote" and slide_spec.quote:
                PptxGenerator._render_quote_layout(slide, slide_spec)
            else:
                PptxGenerator._render_bullets_layout(slide, slide_spec)

        output = io.BytesIO()
        prs.save(output)
        return output.getvalue()

    @staticmethod
    def _render_title_slide(slide: object, title: str, subtitle: str | None) -> None:
        """Render a full-bleed title slide."""
        # Background shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))  # type: ignore[attr-defined]
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.color.rgb = DARK_BG

        # Title & Subtitle text box
        tb = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11.0), Inches(3.2))  # type: ignore[attr-defined]
        tf = tb.text_frame
        tf.word_wrap = True

        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = "Calibri"
        p_title.font.size = Pt(40)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.name = "Calibri"
            p_sub.font.size = Pt(20)
            p_sub.font.color.rgb = MUTED_TEXT
            p_sub.space_before = Pt(14)

    @staticmethod
    def _render_header(slide: object, title: str, subtitle: str | None, slide_num: int) -> None:
        """Render standard top header and slide numbering."""
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.2))  # type: ignore[attr-defined]
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Calibri"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_TEXT

        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.name = "Calibri"
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = SECONDARY_TEXT

        # Slide number
        num_tb = slide.shapes.add_textbox(Inches(12.0), Inches(6.8), Inches(0.8), Inches(0.4))  # type: ignore[attr-defined]
        num_tf = num_tb.text_frame
        p_num = num_tf.paragraphs[0]
        p_num.text = str(slide_num)
        p_num.font.name = "Calibri"
        p_num.font.size = Pt(11)
        p_num.font.color.rgb = MUTED_TEXT
        p_num.alignment = PP_ALIGN.RIGHT

    @staticmethod
    def _render_bullets_layout(slide: object, slide_spec: SlideSpec) -> None:
        """Render bulleted points with clear spacing."""
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(4.8))  # type: ignore[attr-defined]
        tf = tb.text_frame
        tf.word_wrap = True

        for idx, bullet in enumerate(slide_spec.bullets):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.name = "Calibri"
            p.font.size = Pt(16)
            p.font.color.rgb = PRIMARY_TEXT
            p.space_after = Pt(14)

    @staticmethod
    def _render_cards_layout(slide: object, slide_spec: SlideSpec) -> None:
        """Render 2, 3, or 4 visual metric/stat cards."""
        cards = slide_spec.cards
        count = min(len(cards), 4)
        if count == 0:
            return

        total_width = 11.5
        spacing = 0.3
        card_width = (total_width - (spacing * (count - 1))) / count
        card_height = 4.2
        top = 2.0
        left_start = 0.8

        for idx, card in enumerate(cards[:count]):
            c_left = left_start + idx * (card_width + spacing)
            box = slide.shapes.add_shape(  # type: ignore[attr-defined]
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(c_left),
                Inches(top),
                Inches(card_width),
                Inches(card_height),
            )
            box.fill.solid()
            box.fill.fore_color.rgb = CARD_BG
            box.line.color.rgb = CARD_BORDER
            box.line.width = Pt(1)

            tb = slide.shapes.add_textbox(  # type: ignore[attr-defined]
                Inches(c_left + 0.2),
                Inches(top + 0.3),
                Inches(card_width - 0.4),
                Inches(card_height - 0.6),
            )
            tf = tb.text_frame
            tf.word_wrap = True

            # Card Title
            p1 = tf.paragraphs[0]
            p1.text = card.title
            p1.font.name = "Calibri"
            p1.font.size = Pt(13)
            p1.font.bold = True
            p1.font.color.rgb = SECONDARY_TEXT

            # Card Big Value
            p2 = tf.add_paragraph()
            p2.text = card.value
            p2.font.name = "Calibri"
            p2.font.size = Pt(28)
            p2.font.bold = True
            p2.font.color.rgb = ACCENT_BLUE
            p2.space_before = Pt(8)

            # Card Description
            if card.description:
                p3 = tf.add_paragraph()
                p3.text = card.description
                p3.font.name = "Calibri"
                p3.font.size = Pt(11.5)
                p3.font.color.rgb = PRIMARY_TEXT
                p3.space_before = Pt(8)

    @staticmethod
    def _render_comparison_layout(slide: object, slide_spec: SlideSpec) -> None:
        """Render 2-column side-by-side comparison boxes."""
        col_w = Inches(5.6)
        col_h = Inches(4.5)
        top = Inches(1.8)

        # Left Column
        left_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, col_w, col_h)  # type: ignore[attr-defined]
        left_box.fill.solid()
        left_box.fill.fore_color.rgb = CARD_BG
        left_box.line.color.rgb = CARD_BORDER

        ltb = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.0))  # type: ignore[attr-defined]
        ltf = ltb.text_frame
        ltf.word_wrap = True
        for idx, item in enumerate(slide_spec.left_column):
            p = ltf.paragraphs[0] if idx == 0 else ltf.add_paragraph()
            p.text = f"• {item}"
            p.font.name = "Calibri"
            p.font.size = Pt(14)
            p.font.color.rgb = PRIMARY_TEXT
            p.space_after = Pt(10)

        # Right Column
        right_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), top, col_w, col_h)  # type: ignore[attr-defined]
        right_box.fill.solid()
        right_box.fill.fore_color.rgb = CARD_BG
        right_box.line.color.rgb = CARD_BORDER

        rtb = slide.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.2), Inches(4.0))  # type: ignore[attr-defined]
        rtf = rtb.text_frame
        rtf.word_wrap = True
        for idx, item in enumerate(slide_spec.right_column):
            p = rtf.paragraphs[0] if idx == 0 else rtf.add_paragraph()
            p.text = f"• {item}"
            p.font.name = "Calibri"
            p.font.size = Pt(14)
            p.font.color.rgb = PRIMARY_TEXT
            p.space_after = Pt(10)

    @staticmethod
    def _render_quote_layout(slide: object, slide_spec: SlideSpec) -> None:
        """Render a large highlight quote slide."""
        tb = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10.3), Inches(3.5))  # type: ignore[attr-defined]
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"“{slide_spec.quote}”"
        p.font.name = "Calibri"
        p.font.size = Pt(28)
        p.font.italic = True
        p.font.color.rgb = PRIMARY_TEXT

        if slide_spec.author:
            p_auth = tf.add_paragraph()
            p_auth.text = f"— {slide_spec.author}"
            p_auth.font.name = "Calibri"
            p_auth.font.size = Pt(16)
            p_auth.font.bold = True
            p_auth.font.color.rgb = ACCENT_BLUE
            p_auth.space_before = Pt(14)
