"""Professional Word document (.docx) generator using python-docx."""

import base64
import io
from typing import Any

import httpx
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from assistant_core.artifacts.schemas import DocumentSpec, TableSpec

# Theme palettes
DOCX_THEMES: dict[str, dict[str, Any]] = {
    "slate": {
        "primary": RGBColor(30, 41, 59),  # Slate 800
        "secondary": RGBColor(71, 85, 105),  # Slate 600
        "muted": RGBColor(100, 116, 139),  # Slate 500
        "accent_hex": "3B82F6",  # Blue 500
        "header_bg_hex": "1E293B",  # Slate 800
        "zebra_bg_hex": "F8FAFC",  # Slate 50
        "callout_bg_hex": "F1F5F9",  # Slate 100
    },
    "navy": {
        "primary": RGBColor(11, 25, 44),
        "secondary": RGBColor(65, 90, 119),
        "muted": RGBColor(119, 141, 169),
        "accent_hex": "008DDA",
        "header_bg_hex": "0B192C",
        "zebra_bg_hex": "F5F7FA",
        "callout_bg_hex": "EEF2F6",
    },
    "emerald": {
        "primary": RGBColor(6, 78, 59),
        "secondary": RGBColor(22, 101, 52),
        "muted": RGBColor(110, 160, 130),
        "accent_hex": "10B981",
        "header_bg_hex": "064E3B",
        "zebra_bg_hex": "F0FDF4",
        "callout_bg_hex": "ECFDF5",
    },
    "crimson": {
        "primary": RGBColor(76, 5, 25),
        "secondary": RGBColor(159, 18, 57),
        "muted": RGBColor(190, 110, 130),
        "accent_hex": "F43F5E",
        "header_bg_hex": "4C0519",
        "zebra_bg_hex": "FFF1F2",
        "callout_bg_hex": "FFE4E6",
    },
    "dark": {
        "primary": RGBColor(15, 23, 42),
        "secondary": RGBColor(51, 65, 85),
        "muted": RGBColor(100, 116, 139),
        "accent_hex": "60A5FA",
        "header_bg_hex": "0F172A",
        "zebra_bg_hex": "F1F5F9",
        "callout_bg_hex": "E2E8F0",
    },
}


def _resolve_image_stream(image_url: str | None, image_base64: str | None) -> io.BytesIO | None:
    """Fetch and validate an image buffer from URL or base64 data."""
    raw_bytes: bytes | None = None

    if image_base64:
        try:
            b64_str = image_base64
            if "," in b64_str:
                b64_str = b64_str.split(",", 1)[1]
            b64_padded = b64_str + "=" * (-len(b64_str) % 4)
            raw_bytes = base64.b64decode(b64_padded)
        except Exception:  # noqa: BLE001
            raw_bytes = None

    elif image_url and image_url.startswith(("http://", "https://")):
        try:
            with httpx.Client(timeout=8.0, follow_redirects=True) as client:
                res = client.get(image_url)
                if res.status_code == 200:
                    raw_bytes = res.content
        except Exception:  # noqa: BLE001
            raw_bytes = None

    if not raw_bytes:
        return None

    try:
        buf = io.BytesIO(raw_bytes)
        img = Image.open(buf)
        img.verify()
        return io.BytesIO(raw_bytes)
    except Exception:  # noqa: BLE001
        return None


def _set_cell_background(cell: object, hex_color: str) -> None:
    """Set background color of a table cell."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)  # type: ignore[attr-defined]


def _set_cell_margins(
    cell: object, top: int = 120, bottom: int = 120, left: int = 160, right: int = 160
) -> None:
    """Set inner padding for a table cell in twentieths of a point (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()  # type: ignore[attr-defined]
    tcMar = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


class DocxGenerator:
    """Renders structured DocumentSpec into polished .docx bytes."""

    @staticmethod
    def generate(spec: DocumentSpec) -> bytes:
        doc = Document()
        theme = DOCX_THEMES.get(spec.theme, DOCX_THEMES["slate"])

        # Set standard margins (1 inch)
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # 1. Document Title
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(4)
        run_title = title_p.add_run(spec.title)
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(24)
        run_title.font.bold = True
        run_title.font.color.rgb = theme["primary"]

        # Subtitle / Author
        if spec.subtitle:
            sub_p = doc.add_paragraph()
            sub_p.paragraph_format.space_before = Pt(0)
            sub_p.paragraph_format.space_after = Pt(8)
            run_sub = sub_p.add_run(spec.subtitle)
            run_sub.font.name = "Calibri"
            run_sub.font.size = Pt(14)
            run_sub.font.color.rgb = theme["secondary"]

        if spec.author:
            auth_p = doc.add_paragraph()
            auth_p.paragraph_format.space_before = Pt(0)
            auth_p.paragraph_format.space_after = Pt(18)
            run_auth = auth_p.add_run(f"Author: {spec.author}")
            run_auth.font.name = "Calibri"
            run_auth.font.size = Pt(10)
            run_auth.font.italic = True
            run_auth.font.color.rgb = theme["muted"]

        # 2. Sections
        for sec in spec.sections:
            # Heading
            if sec.heading:
                h_p = doc.add_paragraph()
                h_p.paragraph_format.keep_with_next = True
                h_p.paragraph_format.space_before = Pt(16)
                h_p.paragraph_format.space_after = Pt(6)

                h_run = h_p.add_run(sec.heading)
                h_run.font.name = "Calibri"
                h_run.font.bold = True

                if sec.level == 1:
                    h_run.font.size = Pt(16)
                    h_run.font.color.rgb = theme["primary"]
                elif sec.level == 2:
                    h_run.font.size = Pt(13.5)
                    h_run.font.color.rgb = theme["secondary"]
                else:
                    h_run.font.size = Pt(12)
                    h_run.font.color.rgb = theme["secondary"]

            # Paragraphs
            for text in sec.paragraphs:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(text)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.color.rgb = theme["primary"]

            # Bullet points
            for bullet in sec.bullets:
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(3)
                r = bp.add_run(bullet)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.color.rgb = theme["primary"]

            # Image Insertion
            if sec.image_url or sec.image_base64:
                img_stream = _resolve_image_stream(sec.image_url, sec.image_base64)
                if img_stream:
                    try:
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_before = Pt(8)
                        p_img.paragraph_format.space_after = Pt(4)
                        p_img.add_run().add_picture(img_stream, width=Inches(5.5))

                        if sec.image_caption:
                            p_cap = doc.add_paragraph()
                            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p_cap.paragraph_format.space_before = Pt(0)
                            p_cap.paragraph_format.space_after = Pt(8)
                            r_cap = p_cap.add_run(sec.image_caption)
                            r_cap.font.name = "Calibri"
                            r_cap.font.italic = True
                            r_cap.font.size = Pt(9.5)
                            r_cap.font.color.rgb = theme["muted"]
                    except Exception:  # noqa: BLE001, S110
                        pass

            # Callout Note
            if sec.callout:
                callout_tbl = doc.add_table(rows=1, cols=1)
                callout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                callout_cell = callout_tbl.cell(0, 0)
                _set_cell_background(callout_cell, theme["callout_bg_hex"])
                _set_cell_margins(callout_cell, top=140, bottom=140, left=200, right=160)

                tcPr = callout_cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(
                    f"<w:tcBorders {nsdecls('w')}>"
                    f'<w:top w:val="none"/>'
                    f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{theme["accent_hex"]}"/>'
                    f'<w:bottom w:val="none"/>'
                    f'<w:right w:val="none"/>'
                    f"</w:tcBorders>"
                )
                tcPr.append(tcBorders)

                cp = callout_cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(0)
                cr = cp.add_run(f"💡 {sec.callout}")
                cr.font.name = "Calibri"
                cr.font.size = Pt(10.5)
                cr.font.color.rgb = theme["primary"]

                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_before = Pt(6)
                spacer.paragraph_format.space_after = Pt(0)

            # Embedded Table
            if sec.table and sec.table.headers:
                DocxGenerator._render_table(doc, sec.table, theme)

        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()

    @staticmethod
    def _render_table(doc: Any, tbl_spec: TableSpec, theme: dict[str, Any]) -> None:
        """Render a styled data table inside the Word document."""
        num_rows = len(tbl_spec.rows) + 1
        num_cols = len(tbl_spec.headers)
        tbl = doc.add_table(rows=num_rows, cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # 1. Header Row
        for c_idx, h_text in enumerate(tbl_spec.headers):
            cell = tbl.cell(0, c_idx)
            _set_cell_background(cell, theme["header_bg_hex"])
            _set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h_text)
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

        # 2. Data Rows
        for r_idx, row_data in enumerate(tbl_spec.rows):
            fill_hex = theme["zebra_bg_hex"] if r_idx % 2 == 1 else "FFFFFF"
            for c_idx in range(num_cols):
                cell = tbl.cell(r_idx + 1, c_idx)
                _set_cell_background(cell, fill_hex)
                _set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                val_text = str(row_data[c_idx]) if c_idx < len(row_data) else ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(val_text)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = theme["primary"]

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(8)
        spacer.paragraph_format.space_after = Pt(0)
