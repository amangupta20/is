"""Unit tests for XLSX, DOCX, and PPTX artifact generators."""

import base64
import io

import openpyxl
from docx import Document
from PIL import Image
from pptx import Presentation

from assistant_core.artifacts.generators.docx_gen import DocxGenerator
from assistant_core.artifacts.generators.pptx_gen import PptxGenerator
from assistant_core.artifacts.generators.xlsx_gen import XlsxGenerator
from assistant_core.artifacts.schemas import (
    ChartSeries,
    DocumentSectionSpec,
    DocumentSpec,
    PresentationSpec,
    SheetSpec,
    SlideSpec,
    StatCard,
    TableSpec,
    TimelineStep,
    WorkbookSpec,
)


def test_xlsx_generator_creates_valid_workbook() -> None:
    spec = WorkbookSpec(
        title="Q3 Revenue",
        sheets=[
            SheetSpec(
                name="Summary",
                headers=["Department", "Q1", "Q2", "Q3"],
                rows=[
                    ["Engineering", 120000, 135000, 142000],
                    ["Marketing", 45000, 52000, 48000],
                    ["Operations", 30000, 31000, 33000],
                ],
                column_types=["text", "currency", "currency", "currency"],
                totals_row=True,
            )
        ],
    )
    data = XlsxGenerator.generate(spec)
    assert isinstance(data, bytes)
    assert len(data) > 0

    # Load back with openpyxl to verify integrity
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "Summary" in wb.sheetnames
    ws = wb["Summary"]
    assert ws["A1"].value == "Department"
    assert ws["B2"].value == 120000
    assert ws["A5"].value == "Total"
    assert "=SUM(B2:B4)" in str(ws["B5"].value)


def test_docx_generator_creates_valid_document() -> None:
    spec = DocumentSpec(
        title="Architecture Decision Record: Artifact Management",
        subtitle="Durable storage and Office document rendering",
        author="Lead Architect",
        sections=[
            DocumentSectionSpec(
                heading="Context & Problem Statement",
                level=1,
                paragraphs=[
                    "The system requires durable storage for generated documents.",
                    "Files must be editable in OnlyOffice without losing version history.",
                ],
                bullets=[
                    "Sub-millisecond local volume reads",
                    "Immutable version chain",
                ],
                callout="All volumes are backed up to S3 in daily scheduled runs.",
                table=TableSpec(
                    headers=["Format", "Engine", "MIME Type"],
                    rows=[
                        [
                            "XLSX",
                            "openpyxl",
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        ],
                        [
                            "DOCX",
                            "python-docx",
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        ],
                    ],
                ),
            )
        ],
    )
    data = DocxGenerator.generate(spec)
    assert isinstance(data, bytes)
    assert len(data) > 0

    # Load back with python-docx
    doc = Document(io.BytesIO(data))
    paragraphs_text = [p.text for p in doc.paragraphs if p.text]
    assert any("Architecture Decision Record" in t for t in paragraphs_text)
    assert any("Context & Problem Statement" in t for t in paragraphs_text)
    assert len(doc.tables) >= 2  # 1 callout table + 1 data table


def test_pptx_generator_creates_valid_presentation() -> None:
    spec = PresentationSpec(
        title="Homelab Modernization",
        subtitle="Portainer to Komodo Migration",
        slides=[
            SlideSpec(
                title="Migration Milestones",
                layout="bullets",
                bullets=[
                    "Phase 1: Isolated Open Terminal VM setup",
                    "Phase 2: GitOps stacks repository setup",
                    "Phase 3: Komodo REST API integration",
                ],
            ),
            SlideSpec(
                title="Target Metrics",
                layout="cards",
                cards=[
                    StatCard(
                        title="Uptime", value="99.99%", description="Zero downtime deployment"
                    ),
                    StatCard(
                        title="Memory Saved",
                        value="4.2 GB",
                        description="Reduced background overhead",
                    ),
                ],
            ),
            SlideSpec(
                title="Architecture Comparison",
                layout="comparison",
                left_column=["Portainer Web UI", "Manual Stack Edits", "Local State"],
                right_column=["Komodo Orchestrator", "Declarative GitOps", "Immutable Versions"],
            ),
            SlideSpec(
                title="Guiding Principle",
                layout="quote",
                quote="Automate everything, verify continuously, recover instantly.",
                author="SRE Handbook",
            ),
        ],
    )
    data = PptxGenerator.generate(spec)
    assert isinstance(data, bytes)
    assert len(data) > 0

    # Load back with python-pptx
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 5  # 1 Title + 4 Content slides


def test_docx_generator_with_image_and_themes() -> None:
    # Create dummy PNG base64
    img = Image.new("RGB", (60, 60), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    b64_img = base64.b64encode(buf.getvalue()).decode("utf-8")

    spec = DocumentSpec(
        title="Visual Document Test",
        subtitle="Testing Embedded Figures",
        theme="emerald",
        sections=[
            DocumentSectionSpec(
                heading="Architecture Diagram",
                paragraphs=["This section contains an embedded architecture diagram."],
                image_base64=b64_img,
                image_caption="Figure 1: High Level Architecture",
            )
        ],
    )
    data = DocxGenerator.generate(spec)
    assert isinstance(data, bytes)
    assert len(data) > 0

    doc = Document(io.BytesIO(data))
    assert any("Architecture Diagram" in p.text for p in doc.paragraphs)
    assert any("Figure 1: High Level Architecture" in p.text for p in doc.paragraphs)


def test_pptx_generator_with_images_charts_and_timeline() -> None:
    img = Image.new("RGB", (80, 80), color="green")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    b64_img = base64.b64encode(buf.getvalue()).decode("utf-8")

    spec = PresentationSpec(
        title="Visual PPTX Suite",
        theme="navy",
        slides=[
            SlideSpec(
                title="System Overview",
                layout="image_right",
                bullets=["Microservices architecture", "Real-time streaming"],
                image_base64=b64_img,
            ),
            SlideSpec(
                title="Hero Graphic",
                layout="full_image",
                image_base64=b64_img,
                image_caption="High-Resolution Topology",
            ),
            SlideSpec(
                title="Performance Breakdown",
                layout="chart",
                chart_type="column",
                chart_categories=["Q1", "Q2", "Q3", "Q4"],
                chart_series=[ChartSeries(name="Throughput", values=[100, 220, 310, 450])],
            ),
            SlideSpec(
                title="Execution Roadmap",
                layout="timeline",
                timeline_steps=[
                    TimelineStep(step="Phase 1", title="Alpha", description="Core prototype"),
                    TimelineStep(step="Phase 2", title="Beta", description="User onboarding"),
                    TimelineStep(step="Phase 3", title="GA", description="Production launch"),
                ],
            ),
        ],
    )
    data = PptxGenerator.generate(spec)
    assert isinstance(data, bytes)
    assert len(data) > 0

    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 5  # 1 Title + 4 Slides


def test_pdf_generator_creates_valid_pdf() -> None:
    from assistant_core.artifacts.generators.pdf_gen import PdfGenerator

    spec = DocumentSpec(
        title="Architecture Decision Record: PDF Generation",
        subtitle="ReportLab-based PDF generation engine",
        author="Lead Architect",
        theme="slate",
        sections=[
            DocumentSectionSpec(
                heading="Context & Problem Statement",
                level=1,
                paragraphs=[
                    "The assistant requires high quality styled PDF generation.",
                    "PDFs must render cover headers, sections, callouts, and styled tables.",
                ],
                bullets=[
                    "Fast Platypus layout engine",
                    "Dynamic page numbering and running headers",
                    "Zero external browser dependencies",
                ],
                callout="All generated PDFs adhere to standard corporate theme palettes.",
                table=TableSpec(
                    headers=["Module", "Role", "Version"],
                    rows=[
                        ["reportlab", "PDF Engine", "5.0.1"],
                        ["assistant_core", "API Core", "0.1.0"],
                    ],
                ),
            ),
            DocumentSectionSpec(
                heading="Subsection Analysis",
                level=2,
                paragraphs=["Detailed secondary analysis paragraph."],
            ),
        ],
    )
    data = PdfGenerator.generate(spec)
    assert isinstance(data, bytes)
    assert len(data) > 0
    assert data.startswith(b"%PDF-")


def test_pdf_generator_with_image_and_themes() -> None:
    from assistant_core.artifacts.generators.pdf_gen import PdfGenerator

    img = Image.new("RGB", (60, 60), color="red")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    b64_img = base64.b64encode(buf.getvalue()).decode("utf-8")

    for theme in ("slate", "navy", "emerald", "crimson", "dark"):
        spec = DocumentSpec(
            title=f"Theme {theme.title()} PDF Test",
            subtitle="Testing Embedded Figures & Colors",
            theme=theme,  # type: ignore[arg-type]
            sections=[
                DocumentSectionSpec(
                    heading="Visual Evidence",
                    paragraphs=["This section contains an embedded diagram."],
                    image_base64=b64_img,
                    image_caption=f"Figure 1: {theme.title()} Test",
                    callout="Highlight box with theme accent.",
                )
            ],
        )
        data = PdfGenerator.generate(spec)
        assert isinstance(data, bytes)
        assert len(data) > 0
        assert data.startswith(b"%PDF-")
